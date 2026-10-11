#!/usr/bin/env python3
"""
make_report_docx.py -- render docs/TECHNICAL_REPORT.md as a .docx for
submission.

A deliberately small Markdown subset is supported, matching exactly what the
report uses: ATX headings, pipe tables, fenced code blocks, bullet and ordered
lists, blockquotes, horizontal rules, and inline bold / italic / code / links.
A general Markdown engine would be overkill and would drag in dependencies
that are awkward to install on a lab machine.

Output is A4 with Calibri body text, Consolas for code, and a table style that
survives being printed in black and white.

Usage:
    python tools/make_report_docx.py
    python tools/make_report_docx.py --in docs/RESULTS.md --out results.docx
"""

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm
except ImportError:
    sys.exit("python-docx is required:\n    pip install python-docx")

ROOT = Path(__file__).resolve().parent.parent

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x44, 0x44, 0x44)


# --------------------------------------------------------------------------- #
#  low-level helpers
# --------------------------------------------------------------------------- #

def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def code_block_shading(par):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), "F4F5F7")
    par._p.get_or_add_pPr().append(el)


INLINE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\))"
)


def add_inline(par, text):
    """Render bold / italic / code / links inside a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("***") and part.endswith("***"):
            r = par.add_run(part[3:-3]); r.bold = True; r.italic = True
        elif part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xA3, 0x1D, 0x4E)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if m:
                r = par.add_run(m.group(1))
                r.font.color.rgb = ACCENT
                r.underline = True
            else:
                par.add_run(part)
        else:
            par.add_run(part)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|[\s:|-]*$", line)) and "-" in line


# --------------------------------------------------------------------------- #
#  document setup
# --------------------------------------------------------------------------- #

def build_document():
    doc = Document()

    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for attr in ("top_margin", "bottom_margin"):
        setattr(s, attr, Cm(2.0))
    for attr in ("left_margin", "right_margin"):
        setattr(s, attr, Cm(2.2))

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, colour, before in (
        ("Heading 1", 17, ACCENT, 20),
        ("Heading 2", 13.5, ACCENT, 14),
        ("Heading 3", 11.5, GREY, 10),
        ("Heading 4", 10.5, GREY, 8),
    ):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = colour
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    return doc


def add_code(doc, lines):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.4)
    pf.space_before = Pt(6)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    code_block_shading(par)

    for i, ln in enumerate(lines):
        if i:
            par.add_run("\n")
        r = par.add_run(ln)
        r.font.name = CODE_FONT
        r.font.size = Pt(8.5)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, text)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
        shade(cell, "DCE6F1")

    for row in body:
        cells = t.add_row().cells
        for i, text in enumerate(row[:len(header)]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text)
            for r in p.runs:
                r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# --------------------------------------------------------------------------- #
#  the converter
# --------------------------------------------------------------------------- #

def convert(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = build_document()

    i = 0
    first_heading = True

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code(doc, block)
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # horizontal rule
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            doc.add_paragraph().add_run().add_break()
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and first_heading:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                r.bold = True
                r.font.size = Pt(19)
                r.font.color.rgb = ACCENT
                p.paragraph_format.space_after = Pt(4)
                first_heading = False
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 4)}")
                add_inline(p, text)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            add_inline(p, " ".join(x for x in block if x))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = GREY
            continue

        # bullet / ordered list
        mb = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        mo = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if mb or mo:
            indent = len((mb or mo).group(1)) // 2
            text = mb.group(2) if mb else mo.group(3)
            style = "List Bullet" if mb else "List Number"

            # Absorb wrapped continuation lines. Without this, a list item that
            # spans two source lines becomes a list item plus a stray body
            # paragraph, which looks like a formatting error in the .docx.
            i += 1
            while i < len(lines):
                nxt = lines[i]
                stripped_nxt = nxt.strip()
                if (not stripped_nxt
                        or not nxt.startswith((" ", "\t"))
                        or re.match(r"^\s*[-*+]\s+", nxt)
                        or re.match(r"^\s*\d+\.\s+", nxt)
                        or stripped_nxt.startswith(("#", "|", ">", "```"))):
                    break
                text += " " + stripped_nxt
                i += 1

            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
                text = ("• " if mb else "") + text
            p.paragraph_format.left_indent = Cm(0.6 + 0.5 * indent)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text)
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # paragraph: join continuation lines so wrapping is Word's job
        block = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```", "- ", "* ", "+ "))
                    or re.match(r"^\d+\.\s", nxt)
                    or re.match(r"^(-{3,}|\*{3,}|_{3,})$", nxt)):
                break
            block.append(nxt)
            i += 1
        add_inline(doc.add_paragraph(), " ".join(block))

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="src", default="docs/TECHNICAL_REPORT.md")
    ap.add_argument("--out", dest="dst", default=None)
    args = ap.parse_args()

    src = (ROOT / args.src) if not Path(args.src).is_absolute() else Path(args.src)
    if not src.exists():
        sys.exit(f"error: no such file: {src}")

    dst = Path(args.dst) if args.dst else src.with_suffix(".docx")
    if not dst.is_absolute():
        dst = ROOT / dst

    convert(src, dst)
    print(f"Wrote {dst.relative_to(ROOT) if dst.is_relative_to(ROOT) else dst}")
    print(f"      {dst.stat().st_size // 1024} kB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
